import os
import csv
import logging
from typing import List, Dict, Any, Optional
import charset_normalizer
import mistune

from src.domain.entities import ParsedDocument, ParsedSection
from src.domain.ports import DocumentParser
from src.config import settings

logger = logging.getLogger(__name__)

# Utility to format excel grids as Markdown tables for the LLM
def format_grid_as_markdown_table(grid: List[List[str]]) -> str:
    if not grid:
        return ""
    headers = grid[0]
    rows = grid[1:]
    
    header_str = "| " + " | ".join(headers) + " |\n"
    separator_str = "| " + " | ".join(["---"] * len(headers)) + " |\n"
    rows_str = ""
    for row in rows:
        padded_row = row + [""] * (len(headers) - len(row))
        padded_row = padded_row[:len(headers)]
        rows_str += "| " + " | ".join(padded_row) + " |\n"
    return header_str + separator_str + rows_str


class PDFParser(DocumentParser):
    def parse(self, file_path: str, document_id: str) -> ParsedDocument:
        import pdfplumber
        logger.info(f"Iniciando parsing do PDF: {file_path}")
        sections = []
        text_extracted = ""
        is_ocr_needed = False
        
        # Try normal text extraction
        try:
            with pdfplumber.open(file_path) as pdf:
                for i, page in enumerate(pdf.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text_extracted += page_text + "\n"
        except Exception as e:
            logger.warning(f"Falha ao abrir PDF com pdfplumber, forçando OCR: {e}")
            is_ocr_needed = True
            
        if len(text_extracted.strip()) < 50:
            is_ocr_needed = True
            
        if is_ocr_needed:
            logger.info(f"Pouco ou nenhum texto extraível encontrado. Aplicando OCR via Tesseract: {file_path}")
            sections = self._run_ocr(file_path)
        else:
            # Parse structurally
            with pdfplumber.open(file_path) as pdf:
                for i, page in enumerate(pdf.pages):
                    page_text = page.extract_text() or ""
                    page_num = i + 1
                    
                    tables_data = page.extract_tables()
                    formatted_tables = []
                    for table in tables_data:
                        cleaned_table = [[str(cell) if cell is not None else "" for cell in row] for row in table]
                        formatted_tables.append(cleaned_table)
                        
                    sections.append(ParsedSection(
                        heading=f"Página {page_num}",
                        level=3,
                        content=page_text,
                        page_number=page_num,
                        tables=formatted_tables
                    ))
                    
        return ParsedDocument(
            document_id=document_id,
            source_format="PDF",
            title=os.path.basename(file_path),
            sections=sections,
            raw_metadata={"ocr_applied": is_ocr_needed}
        )

    def _run_ocr(self, file_path: str) -> List[ParsedSection]:
        from pdf2image import convert_from_path
        import pytesseract
        
        # Configure tesseract executable path if set in config
        if settings.TESSERACT_CMD:
            pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD
            
        sections = []
        try:
            # Convert PDF pages to images
            poppler_args = {}
            if settings.POPPLER_PATH:
                poppler_args["poppler_path"] = settings.POPPLER_PATH
                
            images = convert_from_path(file_path, **poppler_args)
            for i, image in enumerate(images):
                page_num = i + 1
                ocr_text = pytesseract.image_to_string(image, lang="por")
                
                sections.append(ParsedSection(
                    heading=f"Página {page_num} (OCR)",
                    level=3,
                    content=ocr_text,
                    page_number=page_num,
                    tables=[]
                ))
        except Exception as e:
            logger.error(f"Erro crítico no OCR do PDF: {e}")
            raise RuntimeError(f"Erro ao processar OCR no PDF: {e}")
        return sections


class DOCXParser(DocumentParser):
    def parse(self, file_path: str, document_id: str) -> ParsedDocument:
        import docx
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        
        logger.info(f"Iniciando parsing do DOCX: {file_path}")
        doc = docx.Document(file_path)
        sections = []
        
        current_heading = "Geral"
        current_level = 2
        current_text_blocks = []
        current_tables = []
        
        def flush_section():
            if current_text_blocks or current_tables:
                content = "\n".join(current_text_blocks)
                sections.append(ParsedSection(
                    heading=current_heading,
                    level=current_level,
                    content=content,
                    tables=current_tables.copy()
                ))
                current_text_blocks.clear()
                current_tables.clear()
                
        body_elements = doc.element.body
        for child in body_elements.iterchildren():
            tag = child.tag
            if tag.endswith('p'):
                p = Paragraph(child, doc)
                text = p.text.strip()
                if not text:
                    continue
                    
                if p.style.name.startswith('Heading'):
                    flush_section()
                    current_heading = text
                    try:
                        current_level = int(p.style.name.split()[-1])
                    except Exception:
                        current_level = 2
                else:
                    current_text_blocks.append(text)
                    
            elif tag.endswith('tbl'):
                t = Table(child, doc)
                grid = []
                for row in t.rows:
                    row_vals = [cell.text.strip() for cell in row.cells]
                    grid.append(row_vals)
                
                if grid:
                    current_tables.append(grid)
                    tbl_md = format_grid_as_markdown_table(grid)
                    current_text_blocks.append(tbl_md)
                    
        flush_section()
        
        return ParsedDocument(
            document_id=document_id,
            source_format="DOCX",
            title=os.path.basename(file_path),
            sections=sections
        )


class XLSXParser(DocumentParser):
    def parse(self, file_path: str, document_id: str) -> ParsedDocument:
        import openpyxl
        logger.info(f"Iniciando parsing do XLSX (documento lógico único): {file_path}")
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sections = []
        
        try:
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                grid = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(val) if val is not None else "" for val in row]
                    if any(row_vals):
                        grid.append(row_vals)
                        
                content = ""
                if grid:
                    content = format_grid_as_markdown_table(grid)
                    
                sections.append(ParsedSection(
                    heading=sheet_name,
                    level=2,
                    content=content,
                    tables=[grid] if grid else []
                ))
        finally:
            wb.close()
            
        return ParsedDocument(
            document_id=document_id,
            source_format="XLSX",
            title=os.path.basename(file_path),
            sections=sections
        )


class CSVParser(DocumentParser):
    def parse(self, file_path: str, document_id: str) -> ParsedDocument:
        logger.info(f"Iniciando parsing do CSV: {file_path}")
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            grid = [row for row in reader if row]
            
        content = format_grid_as_markdown_table(grid) if grid else ""
        
        return ParsedDocument(
            document_id=document_id,
            source_format="CSV",
            title=os.path.basename(file_path),
            sections=[ParsedSection(
                heading="Dados CSV",
                level=2,
                content=content,
                tables=[grid] if grid else []
            )]
        )


class TXTParser(DocumentParser):
    def parse(self, file_path: str, document_id: str) -> ParsedDocument:
        logger.info(f"Iniciando parsing do TXT: {file_path}")
        with open(file_path, "rb") as f:
            raw_data = f.read()
            
        result = charset_normalizer.detect(raw_data)
        encoding = result['encoding'] or 'utf-8'
        text = raw_data.decode(encoding, errors='ignore')
        
        return ParsedDocument(
            document_id=document_id,
            source_format="TXT",
            title=os.path.basename(file_path),
            sections=[ParsedSection(
                heading="Conteúdo Principal",
                level=2,
                content=text,
                tables=[]
            )]
        )


class MDParser(DocumentParser):
    def parse(self, file_path: str, document_id: str) -> ParsedDocument:
        logger.info(f"Iniciando parsing do MD: {file_path}")
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            lines = f.readlines()
            
        sections = []
        current_heading = "Geral"
        current_level = 2
        current_lines = []
        
        def flush_section():
            if current_lines:
                content = "".join(current_lines).strip()
                sections.append(ParsedSection(
                    heading=current_heading,
                    level=current_level,
                    content=content,
                    tables=[]
                ))
                current_lines.clear()
                
        for line in lines:
            stripped = line.strip()
            if stripped.startswith("#"):
                parts = stripped.split(maxsplit=1)
                hashes = parts[0]
                if all(char == '#' for char in hashes) and len(parts) > 1:
                    flush_section()
                    current_level = len(hashes)
                    current_heading = parts[1]
                    continue
            current_lines.append(line)
            
        flush_section()
        
        return ParsedDocument(
            document_id=document_id,
            source_format="MD",
            title=os.path.basename(file_path),
            sections=sections
        )


class ParserRegistry:
    def __init__(self):
        self._parsers: Dict[str, DocumentParser] = {
            ".pdf": PDFParser(),
            ".docx": DOCXParser(),
            ".xlsx": XLSXParser(),
            ".csv": CSVParser(),
            ".txt": TXTParser(),
            ".md": MDParser()
        }
        
    def get_parser(self, extension: str) -> DocumentParser:
        ext = extension.lower()
        if ext not in self._parsers:
            raise ValueError(f"Formato de arquivo não suportado: {ext}")
        return self._parsers[ext]
        
    def register_parser(self, extension: str, parser: DocumentParser) -> None:
        self._parsers[extension.lower()] = parser
