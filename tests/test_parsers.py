import os
import tempfile
import pytest
from unittest.mock import MagicMock, patch

from src.domain.entities import ParsedDocument, ParsedSection
from src.infrastructure.adapters.parsers import (
    TXTParser, MDParser, CSVParser, XLSXParser, DOCXParser, PDFParser, ParserRegistry
)
from src.infrastructure.adapters.normalizer import RegexTextNormalizer

def test_registry():
    registry = ParserRegistry()
    assert isinstance(registry.get_parser(".txt"), TXTParser)
    assert isinstance(registry.get_parser(".md"), MDParser)
    assert isinstance(registry.get_parser(".csv"), CSVParser)
    assert isinstance(registry.get_parser(".xlsx"), XLSXParser)
    assert isinstance(registry.get_parser(".docx"), DOCXParser)
    assert isinstance(registry.get_parser(".pdf"), PDFParser)
    
    with pytest.raises(ValueError):
        registry.get_parser(".unknown")


def test_txt_parser():
    parser = TXTParser()
    content = "Olá, mundo! Este é um teste com acentuação muito mais longo para garantir que a detecção de encoding funcione corretamente e de forma estável."
    
    # On Windows, close file handle immediately so other processes can access it
    tmp = tempfile.NamedTemporaryFile(suffix=".txt", delete=False)
    tmp_path = tmp.name
    tmp.close()
    
    try:
        with open(tmp_path, "wb") as f:
            f.write(content.encode("utf-8"))
            
        parsed = parser.parse(tmp_path, "doc-123")
        assert parsed.document_id == "doc-123"
        assert parsed.source_format == "TXT"
        assert len(parsed.sections) == 1
        assert parsed.sections[0].heading == "Conteúdo Principal"
        assert parsed.sections[0].content.strip() == content
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


def test_md_parser():
    parser = MDParser()
    content = """# Título Documento
Introdução ao documento.

## Seção 1
Texto da seção 1.

### Subseção 1.1
Texto da subseção 1.1.
"""
    tmp = tempfile.NamedTemporaryFile(suffix=".md", delete=False)
    tmp_path = tmp.name
    tmp.close()
    
    try:
        with open(tmp_path, "w", encoding="utf-8") as f:
            f.write(content)
            
        parsed = parser.parse(tmp_path, "doc-md")
        assert parsed.document_id == "doc-md"
        assert parsed.source_format == "MD"
        assert len(parsed.sections) == 3
        
        # Section 1 (h1)
        assert parsed.sections[0].heading == "Título Documento"
        assert parsed.sections[0].level == 1
        assert "Introdução" in parsed.sections[0].content
        
        # Section 2 (h2)
        assert parsed.sections[1].heading == "Seção 1"
        assert parsed.sections[1].level == 2
        
        # Section 3 (h3)
        assert parsed.sections[2].heading == "Subseção 1.1"
        assert parsed.sections[2].level == 3
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


def test_csv_parser():
    parser = CSVParser()
    content = "Header1,Header2\nValue1,Value2\nValue3,Value4\n"
    
    tmp = tempfile.NamedTemporaryFile(suffix=".csv", delete=False)
    tmp_path = tmp.name
    tmp.close()
    
    try:
        with open(tmp_path, "w", encoding="utf-8") as f:
            f.write(content)
            
        parsed = parser.parse(tmp_path, "doc-csv")
        assert parsed.source_format == "CSV"
        assert len(parsed.sections) == 1
        # Check markdown table conversion
        assert "| Header1 | Header2 |" in parsed.sections[0].content
        assert "| Value1 | Value2 |" in parsed.sections[0].content
        assert len(parsed.sections[0].tables) == 1
        assert parsed.sections[0].tables[0] == [["Header1", "Header2"], ["Value1", "Value2"], ["Value3", "Value4"]]
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


def test_xlsx_parser():
    import openpyxl
    parser = XLSXParser()
    
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Planilha A"
    ws1.append(["ColA", "ColB"])
    ws1.append(["1", "2"])
    
    ws2 = wb.create_sheet("Planilha B")
    ws2.append(["ColX", "ColY"])
    ws2.append(["A", "B"])
    
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp_path = tmp.name
    tmp.close()
    
    try:
        wb.save(tmp_path)
        
        parsed = parser.parse(tmp_path, "doc-xlsx")
        assert parsed.source_format == "XLSX"
        assert len(parsed.sections) == 2
        
        # Check Sheet 1 (Planilha A)
        assert parsed.sections[0].heading == "Planilha A"
        assert "| ColA | ColB |" in parsed.sections[0].content
        assert "| 1 | 2 |" in parsed.sections[0].content
        
        # Check Sheet 2 (Planilha B)
        assert parsed.sections[1].heading == "Planilha B"
        assert "| ColX | ColY |" in parsed.sections[1].content
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


def test_docx_parser():
    import docx
    parser = DOCXParser()
    
    doc = docx.Document()
    doc.add_heading("Título Principal", level=1)
    doc.add_paragraph("Introdução ao Word.")
    doc.add_heading("Subseção Word", level=2)
    doc.add_paragraph("Parágrafo da subseção.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "H1"
    table.cell(0, 1).text = "H2"
    table.cell(1, 0).text = "V1"
    table.cell(1, 1).text = "V2"
    
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp_path = tmp.name
    tmp.close()
    
    try:
        doc.save(tmp_path)
        
        parsed = parser.parse(tmp_path, "doc-docx")
        assert parsed.source_format == "DOCX"
        assert len(parsed.sections) == 2
        
        assert parsed.sections[0].heading == "Título Principal"
        assert parsed.sections[0].level == 1
        assert "Introdução" in parsed.sections[0].content
        
        assert parsed.sections[1].heading == "Subseção Word"
        assert parsed.sections[1].level == 2
        assert "| H1 | H2 |" in parsed.sections[1].content
        assert "| V1 | V2 |" in parsed.sections[1].content
        assert len(parsed.sections[1].tables) == 1
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


@patch("pdfplumber.open")
def test_pdf_parser_text_based(mock_pdfplumber_open):
    # Mocking text based PDF extraction (making string > 50 characters to avoid OCR trigger)
    mock_pdf = MagicMock()
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Conteúdo textual da página do PDF muito mais longo para evitar cair no OCR fallback de menos de 50 caracteres."
    mock_page.extract_tables.return_value = [[["Col1", "Col2"], ["Val1", "Val2"]]]
    mock_pdf.pages = [mock_page]
    mock_pdfplumber_open.return_value.__enter__.return_value = mock_pdf
    
    parser = PDFParser()
    parsed = parser.parse("fake_path.pdf", "doc-pdf")
    
    assert parsed.source_format == "PDF"
    assert len(parsed.sections) == 1
    assert parsed.sections[0].heading == "Página 1"
    assert parsed.sections[0].page_number == 1
    assert "Conteúdo textual" in parsed.sections[0].content
    assert parsed.sections[0].tables == [[["Col1", "Col2"], ["Val1", "Val2"]]]
    assert parsed.raw_metadata["ocr_applied"] is False


@patch("pdfplumber.open")
@patch("pdf2image.convert_from_path")
@patch("pytesseract.image_to_string")
def test_pdf_parser_scanned_ocr(mock_image_to_string, mock_convert, mock_pdfplumber_open):
    # Setup pdfplumber to return no/minimal text
    mock_pdf = MagicMock()
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "" # No text
    mock_page.extract_tables.return_value = []
    mock_pdf.pages = [mock_page]
    mock_pdfplumber_open.return_value.__enter__.return_value = mock_pdf
    
    # Setup OCR mocks
    mock_convert.return_value = [MagicMock()] # 1 page
    mock_image_to_string.return_value = "Texto reconhecido via OCR Tesseract."
    
    parser = PDFParser()
    parsed = parser.parse("fake_path.pdf", "doc-pdf-scanned")
    
    assert parsed.source_format == "PDF"
    assert len(parsed.sections) == 1
    assert parsed.sections[0].heading == "Página 1 (OCR)"
    assert parsed.sections[0].page_number == 1
    assert "reconhecido via OCR" in parsed.sections[0].content
    assert parsed.raw_metadata["ocr_applied"] is True


def test_regex_normalizer():
    normalizer = RegexTextNormalizer()
    
    # Spacing and control characters
    text = "Este   é\t um  teste \x07 com caracteres \x1f de controle."
    clean = normalizer.normalize(text, is_ocr=False)
    assert clean == "Este é um teste com caracteres de controle."
    
    # Redundant blank lines
    text_newlines = "Linha 1\n\n\n\nLinha 2"
    clean_newlines = normalizer.normalize(text_newlines, is_ocr=False)
    assert clean_newlines == "Linha 1\n\nLinha 2"
    
    # Broken hyphenation from PDF line wrap
    text_hyphen = "Isso é uma infor-\nmação importante."
    clean_hyphen = normalizer.normalize(text_hyphen, is_ocr=False)
    assert clean_hyphen == "Isso é uma informação importante."
    
    # OCR specific digit corrections
    ocr_confused = "O valor é R$ 1.O00, e temos 1O% de desconto (l0 reais ou 3I2 reais)."
    clean_ocr = normalizer.normalize(ocr_confused, is_ocr=True)
    # "1.O00" -> "1.000", "1O%" -> "10%", "l0" -> "10", "3I2" -> "312"
    assert "1.000" in clean_ocr
    assert "10%" in clean_ocr
    assert "312" in clean_ocr
    assert "10" in clean_ocr
